"""Пример: python cli.py "Барбершоп, Автомойка" "Уфа" --max 40 --max-rating 4.3 --min-reviews 5"""
import argparse
import asyncio
import math
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter

from parser import search_async

FONT = "Arial"
INK, ACCENT, WARM, COLD = "1B1A2E", "FF6B4A", "FFB020", "4CC3FF"
SOFT, MUTED, LINE = "F4F1FA", "6B6880", "E4E0EF"
BADGE = {"Горячий": "🔥 Горячий", "Тёплый": "🟠 Тёплый", "Холодный": "🔵 Холодный"}
BADGE_COLOR = {"Горячий": ACCENT, "Тёплый": WARM, "Холодный": COLD}


def score(r):
    """Приоритет лида 0-100: чем ниже рейтинг и больше оценок, тем больнее проблема."""
    pain = max(0.0, min(1.0, (4.6 - r["rating"]) / 2.0)) * 60
    volume = min(math.log10((r["reviews"] or 0) + 1) / 2, 1.0) * 25
    contact = 10 if r["phone"] else 0
    no_site = 5 if not r["site"] else 0
    return round(pain + volume + contact + no_site)


def priority(s):
    return "Горячий" if s >= 60 else "Тёплый" if s >= 40 else "Холодный"


def badge(s):
    return BADGE[priority(s)]


def message(r):
    return (
        f"Здравствуйте! Меня зовут [Имя]. Посмотрел карточку «{r['name']}» на Яндекс.Картах: "
        f"рейтинг {str(r['rating']).replace('.', ',')} при {r['reviews']} оценках. "
        "Такой рейтинг обычно отталкивает часть новых клиентов. Мы помогаем поднять его за счёт "
        "настоящих положительных отзывов: настраиваем так, чтобы довольные клиенты сами "
        "оставляли оценку (QR-код, ссылка, короткая инструкция для администраторов). "
        "Рейтинг растёт за счёт потока свежих хороших оценок, а вам не нужно спорить с "
        "недовольными. Могу бесплатно прислать короткий разбор вашей карточки. Интересно?"
    )


def fill(color):
    return PatternFill("solid", fgColor=color)


def rating_color(v):
    return "F8B4A8" if v <= 3.0 else "FDD9A0" if v <= 3.8 else "FFF0B8"


def header(ws, widths, row=1, color=INK):
    for c in ws[row]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
        c.fill = fill(color)
        c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True, indent=1)
    ws.row_dimensions[row].height = 30
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(widths))}{ws.max_row}"


def body(ws, first, height):
    line = Side(style="thin", color=LINE)
    for row in ws.iter_rows(min_row=first):
        ws.row_dimensions[row[0].row].height = height
        for c in row:
            c.font = Font(name=FONT, size=10, color=INK)
            c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
            c.border = Border(bottom=line)


def link(cell, url, text=None):
    cell.value = text or cell.value
    cell.hyperlink = url
    cell.font = Font(name=FONT, size=10, color="2563EB", underline="single")


def summary_sheet(ws, rows, leads):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 14
    ws["B2"] = "Лиды для услуг по отзывам"
    ws["B2"].font = Font(name=FONT, size=22, bold=True, color=INK)
    ws["B3"] = f"Сформировано {datetime.now():%d.%m.%Y в %H:%M}"
    ws["B3"].font = Font(name=FONT, size=10, color=MUTED)
    hot = sum(1 for r in leads if r["score"] >= 60)
    cards = [("Просмотрено", len(rows), INK), ("Подходящих лидов", len(leads), ACCENT),
             ("🔥 Горячих", hot, ACCENT), ("С телефоном", sum(1 for r in leads if r["phone"]), INK)]
    for k, (label, val, color) in enumerate(cards):
        c1, c2 = get_column_letter(2 + 2 * k), get_column_letter(3 + 2 * k)
        ws.merge_cells(f"{c1}5:{c2}5")
        ws.merge_cells(f"{c1}6:{c2}7")
        for r_ in (5, 6, 7):
            for c in (c1, c2):
                ws[f"{c}{r_}"].fill = fill(SOFT)
        ws[f"{c1}5"] = label
        ws[f"{c1}5"].font = Font(name=FONT, size=10, color=MUTED)
        ws[f"{c1}5"].alignment = Alignment(horizontal="left", indent=1, vertical="center")
        ws[f"{c1}6"] = val
        ws[f"{c1}6"].font = Font(name=FONT, size=28, bold=True, color=color)
        ws[f"{c1}6"].alignment = Alignment(horizontal="left", indent=1, vertical="center")
    ws.row_dimensions[5].height = 22

    ws["B10"] = "Кому писать первым"
    ws["B10"].font = Font(name=FONT, size=13, bold=True, color=INK)
    heads = [("B", "D", "Название"), ("E", "E", "Рейтинг"), ("F", "F", "Оценок"),
             ("G", "H", "Телефон"), ("I", "I", "Приоритет")]
    for a, b, t in heads:
        if a != b:
            ws.merge_cells(f"{a}11:{b}11")
        ws[f"{a}11"] = t
        for col in range(ord(a), ord(b) + 1):
            ws[f"{chr(col)}11"].fill = fill(INK)
        ws[f"{a}11"].font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        ws[f"{a}11"].alignment = Alignment(indent=1, vertical="center")
    for i, r in enumerate(leads[:5]):
        n = 12 + i
        vals = {"B": r["name"], "E": f"{r['rating']} ★", "F": r["reviews"],
                "G": r["phone"] or "—", "I": BADGE[priority(r["score"])]}
        for a, b, _ in heads:
            if a != b:
                ws.merge_cells(f"{a}{n}:{b}{n}")
            ws[f"{a}{n}"] = vals[a]
            ws[f"{a}{n}"].font = Font(name=FONT, size=10, color=INK)
            ws[f"{a}{n}"].alignment = Alignment(horizontal="left", indent=1, vertical="center")
            for col in range(ord(a), ord(b) + 1):
                ws[f"{chr(col)}{n}"].border = Border(bottom=Side(style="thin", color=LINE))
        ws.row_dimensions[n].height = 24
    if not leads:
        ws["B12"] = "Подходящих лидов не нашлось: попробуйте поднять порог рейтинга."
        ws["B12"].font = Font(name=FONT, size=10, color=MUTED)


def build_xlsx(rows, out):
    """Собирает Excel (листы «Сводка», «Лиды», «Все организации»), возвращает лиды по убыванию балла."""
    leads = [r for r in rows if r["is_lead"]]
    for r in leads:
        r["score"] = score(r)
    leads.sort(key=lambda r: r["score"], reverse=True)

    wb = Workbook()
    summary_sheet(wb.active, rows, leads)
    wb.active.title = "Сводка"

    ws = wb.create_sheet("Лиды")
    ws.append(["Приоритет", "Балл", "Название", "Рейтинг", "Оценок", "Телефон", "Сайт",
               "Адрес", "Сфера", "Город", "Карточка", "Первое сообщение"])
    for r in leads:
        ws.append([BADGE[priority(r["score"])], r["score"], r["name"], f"{r['rating']} ★",
                   r["reviews"], r["phone"] or "", r["site"] or "нет сайта", r["address"],
                   r["category"], r["city"], "Открыть" if r["url"] else "", message(r)])
    body(ws, 2, 92)
    for i, r in enumerate(leads, 2):
        pr = priority(r["score"])
        c = ws.cell(row=i, column=1)
        c.fill = fill(BADGE_COLOR[pr])
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF" if pr == "Горячий" else INK)
        ws.cell(row=i, column=3).font = Font(name=FONT, size=11, bold=True, color=INK)
        rc = ws.cell(row=i, column=4)
        rc.fill = fill(rating_color(r["rating"]))
        rc.font = Font(name=FONT, size=11, bold=True, color=INK)
        if r["phone"]:
            link(ws.cell(row=i, column=6), "tel:" + re.sub(r"[^\d+]", "", r["phone"]))
        if r["site"] and r["site"].startswith("http"):
            link(ws.cell(row=i, column=7), r["site"])
        if r["url"]:
            link(ws.cell(row=i, column=11), r["url"])
        ws.cell(row=i, column=12).font = Font(name=FONT, size=9, color=MUTED)
        for col in (2, 4, 5):
            ws.cell(row=i, column=col).alignment = Alignment(horizontal="center", vertical="center")
    header(ws, [14, 8, 32, 11, 9, 18, 22, 34, 16, 14, 11, 80])

    ws2 = wb.create_sheet("Все организации")
    ws2.append(["Название", "Рейтинг", "Оценок", "Лид", "Адрес", "Сфера", "Город"])
    for r in rows:
        ws2.append([r["name"], f"{r['rating']} ★" if r["rating"] else "—", r["reviews"],
                    "да" if r["is_lead"] else "нет", r["address"], r["category"], r["city"]])
    body(ws2, 2, 26)
    for i, r in enumerate(rows, 2):
        if r["rating"]:
            ws2.cell(row=i, column=2).fill = fill(rating_color(r["rating"]))
        if r["is_lead"]:
            ws2.cell(row=i, column=4).font = Font(name=FONT, size=10, bold=True, color=ACCENT)
    header(ws2, [34, 11, 9, 7, 40, 18, 14])

    wb.save(out)
    return leads


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("categories", help="сферы через запятую")
    ap.add_argument("cities", help="города через запятую")
    ap.add_argument("--max", type=int, default=40)
    ap.add_argument("--max-rating", type=float, default=4.3)
    ap.add_argument("--min-reviews", type=int, default=5)
    ap.add_argument("--out", default="result.xlsx")
    a = ap.parse_args()

    cats = [x.strip() for x in a.categories.split(",") if x.strip()]
    cities = [x.strip() for x in a.cities.split(",") if x.strip()]

    rows = []
    for city in cities:
        for cat in cats:
            rows += asyncio.run(search_async(
                cat, city, max_results=a.max, max_rating=a.max_rating,
                min_reviews=a.min_reviews))

    leads = build_xlsx(rows, a.out)
    print(f"Всего организаций: {len(rows)}, лидов: {len(leads)}. Файл: {a.out}")


if __name__ == "__main__":
    main()
